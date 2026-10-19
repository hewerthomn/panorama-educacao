"""Fixtures fictícias com os cabeçalhos inspecionados do Inep; nunca publicadas."""
import csv
import json
from pathlib import Path
from zipfile import ZipFile
import pytest
from openpyxl import Workbook
from edu_panorama.ideb.parser import parse, code, numeric
from edu_panorama.ideb.sources import locate_source, validate_file
from edu_panorama.ideb.catalog import SOURCE_FILES, STAGES
from edu_panorama.ideb.etl import generate, state_status
from edu_panorama.cli import main


def workbook(path, level='school', *, value=5.7, duplicate=None, uf='RO'):
    w = Workbook()
    w.remove(w.active)
    stages = list(STAGES) if level == 'state' else ['anos_iniciais']
    for stage in stages:
        sh = w.create_sheet({'anos_iniciais':'UF e Regiões (AI)', 'anos_finais':'UF e Regiões (AF)', 'ensino_medio':'UF e Regiões (EM)'}[stage])
        identity = {'school':['SG_UF','CO_MUNICIPIO','NO_MUNICIPIO','ID_ESCOLA','NO_ESCOLA','REDE'],
                    'municipality':['SG_UF','CO_MUNICIPIO','NO_MUNICIPIO','REDE'], 'state':[None,None]}[level]
        visual = ['Região/\nUnidade da Federação', 'Rede'] if level == 'state' else identity.copy()
        technical = identity.copy()
        sub = [None] * len(identity)
        for year in (2021,2023,2025):
            visual += [f'Taxa de Aprovação - {year}', f'Nota SAEB - {year}', None, None, f'IDEB\n{year}\n(N x P)']
            # Erro de rótulo real das UFs: P/2025 é chamado 2023.
            pyear = 2023 if level == 'state' and year == 2025 else year
            technical += [f'VL_INDICADOR_REND_{pyear}', f'VL_NOTA_MEDIA_{year}', f'VL_NOTA_MATEMATICA_{year}', f'VL_NOTA_PORTUGUES_{year}', f'VL_OBSERVADO_{year}']
            sub += ['Indicador de Rendimento (P)', 'Nota Média Padronizada (N)', 'Matemática','Língua Portuguesa',None]
        technical += ['VL_PROJECAO_2021']
        sh.append(['Ministério da Educação'])
        sh.append(visual); sh.append(sub); sh.append(technical)
        row = {'state':['Rondônia' if uf == 'RO' else 'Acre','Pública'], 'municipality':[uf,1100015.0,'Município fictício','Pública'],
               'school':[uf,1100015.0,'Município fictício',11000001.0,'Escola fictícia','Estadual']}[level]
        for i in range(3): row += [0.91+i*.01, 6.123456, 220.1, '210,25', value if i != 1 else 'ND***']
        row += [9.9]
        sh.append(row)
        if duplicate:
            copy = row.copy()
            if duplicate == 'conflict': copy[-2] = 8
            sh.append(copy)
    w.save(path)
    return path


@pytest.mark.parametrize('value,expected', [(1100015.0,'1100015'),('1100015.0','1100015'),('0012345','0012345')])
def test_codes(value,expected): assert code(value,7) == expected


@pytest.mark.parametrize('value',[1.4,'a','123456789',None])
def test_invalid_codes(value):
    with pytest.raises(ValueError): code(value,7)


@pytest.mark.parametrize('token,status',[('ND','not_published'),('ND*','not_published'),('ND***','not_published'),('-','not_published'),(None,'missing'),('N/A','not_applicable')])
def test_absence(token,status): assert numeric(token,ideb=True) == (None,status)


def test_values_precision():
    assert numeric('5,712345',ideb=True) == (5.712345,'available')
    assert numeric(0,ideb=True) == (0,'available')
    for invalid in ('nan',11,'foo'):
        with pytest.raises(ValueError): numeric(invalid,ideb=True)


@pytest.mark.parametrize('level',['state','municipality','school'])
def test_real_header_structure_multiple_editions_and_levels(tmp_path,level):
    path=workbook(tmp_path/'fixture.xlsx',level)
    records,mapping=parse(path,level=level,stage=None if level=='state' else 'anos_iniciais')
    assert {r['year'] for r in records} == {2021,2023,2025}
    assert {r['level'] for r in records} == {level}
    assert {r['uf_code'] for r in records} == {'11'}
    first=records[0]
    assert first['ideb'] == 5.7
    assert records[1]['ideb'] is None and records[1]['raw_value']=='ND***'
    assert first['components']['performance'] == 6.123456
    assert records[2]['components']['yield'] == .93
    assert first['school_id'] == ('11000001' if level=='school' else None)
    assert len(mapping) == (3 if level=='state' else 1)
    if level=='state': assert {r['stage'] for r in records} == set(STAGES)
    assert parse(path,level=level,stage=None if level=='state' else 'anos_iniciais',years=[2023])[0][0]['ideb'] is None
    with pytest.raises(ValueError,match='edições ausentes'): parse(path,level=level,stage='anos_iniciais',years=[2019])
    with pytest.raises(ValueError,match='Nenhuma'): parse(path,level=level,stage='anos_iniciais',ufs={'AC'})


def test_duplicate_and_zip(tmp_path):
    path=workbook(tmp_path/'fixture.xlsx',duplicate='equal')
    archive=tmp_path/'official.zip'
    with ZipFile(archive,'w') as z:
        z.write(path,'data/fixture.xlsx');z.writestr('data/fixture.ods','duplicate format')
    records,_=parse(archive,level='school',stage='anos_iniciais')
    assert len(records)==3
    workbook(path,duplicate='conflict')
    with pytest.raises(ValueError,match='duplicados'):parse(path,level='school',stage='anos_iniciais')


def test_csv_and_stages(tmp_path):
    path=tmp_path/'fixture.csv'
    path.write_text('SG_UF;CO_MUNICIPIO;NO_MUNICIPIO;ID_ESCOLA;NO_ESCOLA;REDE;VL_OBSERVADO_2021;VL_OBSERVADO_2023;VL_OBSERVADO_2025\nRO;1100015;Cidade;11000001;Escola;Estadual;5,1;-;6,2\n')
    for stage in STAGES:
        records,_=parse(path,level='school',stage=stage)
        assert records[0]['stage']==stage and records[0]['ideb']==5.1
        assert records[1]['status']=='not_published'


def test_integrity_local_cache_download_errors(tmp_path,monkeypatch):
    key='municipios_anos_iniciais'
    path=workbook(tmp_path/'official.xlsx','municipality')
    assert locate_source(key,tmp_path,local_files={key:path})[0]==path
    invalid=tmp_path/'bad.zip';invalid.write_text('<html>error</html>')
    with pytest.raises(ValueError): validate_file(invalid)
    with pytest.raises(FileNotFoundError):locate_source(key,tmp_path,no_download=True)
    with pytest.raises(ValueError):locate_source(key,tmp_path,urls={key:'http://example.com/x.zip'})
    import requests
    calls=[]
    def failed(*args,**kwargs):
        calls.append(kwargs)
        raise requests.exceptions.SSLError('invalid certificate')
    monkeypatch.setattr(requests,'get',failed)
    with pytest.raises(RuntimeError,match='certificados'):locate_source(key,tmp_path)
    assert 'verify' not in calls[0] or calls[0]['verify'] is True
    assert not list(tmp_path.glob('*.part'))


def test_etl_cli_status_skip_force_schema_and_separation(tmp_path,monkeypatch):
    files={}
    for key,(level,stage,_) in SOURCE_FILES.items():files[key]=str(workbook(tmp_path/f'{key}.xlsx',level))
    output=tmp_path/'output'
    settings={'local_files':files}
    generate(['RO'],[2021,2023,2025],output_dir=output,cache_dir=tmp_path,settings=settings,no_download=True)
    assert state_status('RO',[2021,2023,2025],output)==(True,'completo')
    manifest=json.loads((output/'RO/manifest.json').read_text())
    for stage in STAGES:
        for level in ('state','municipality','school'):
            data=json.loads((output/'RO'/manifest['files'][stage][level]['path']).read_text())
            assert {r['level'] for r in data['records']} == {level}
            assert {r['stage'] for r in data['records']} == {stage}
    import edu_panorama.ideb.etl as etl
    original=etl.parse
    def fail(*args,**kwargs):raise AssertionError('cache must skip parsing')
    monkeypatch.setattr(etl,'parse',fail)
    generate(['RO'],[2021,2023,2025],output_dir=output,cache_dir=tmp_path,settings=settings,no_download=True)
    monkeypatch.setattr(etl,'parse',original)
    generate(['RO'],[2021,2023,2025],output_dir=output,cache_dir=tmp_path,settings=settings,no_download=True,force=True)
    path=output/'RO'/manifest['files']['anos_iniciais']['school']['path'];path.write_text('{}')
    assert not state_status('RO',[2021],output)[0]
    manifest['schema_version']=0;(output/'RO/manifest.json').write_text(json.dumps(manifest))
    assert not state_status('RO',[2021],output)[0]
    config=tmp_path/'config.toml';config.write_text('[ideb]\noutput_dir="output"\n')
    assert main(['ideb','--config',str(config),'status','--uf','RO'])==0
    assert main(['ideb','--config',str(config),'generate','--uf','RO','--no-download'])==1


@pytest.mark.parametrize('name,uf', [('R. G. do Norte','RN'), ('R. G. do Sul','RS'), ('M. G. do Sul','MS')])
def test_official_state_abbreviations_and_network_footnotes(tmp_path,name,uf):
    path=workbook(tmp_path/'fixture.xlsx','state')
    from openpyxl import load_workbook
    w=load_workbook(path)
    for sheet in w:
        sheet.cell(5,1,name)
        sheet.cell(5,2,'Pública (4)')
    w.save(path)
    records,_=parse(path,level='state',ufs={uf})
    assert {r['uf'] for r in records} == {uf}
    assert {r['network'] for r in records} == {'Pública'}
    assert {r['network_label'] for r in records} == {'Pública (4)'}


def test_successful_download_and_cache_are_verified(tmp_path,monkeypatch):
    from io import BytesIO
    path=workbook(tmp_path/'fixture.xlsx','municipality')
    stream=BytesIO()
    with ZipFile(stream,'w') as archive: archive.write(path,'official.xlsx')
    payload=stream.getvalue()
    calls=[]
    class Response:
        def __enter__(self): return self
        def __exit__(self,*args): pass
        def raise_for_status(self): pass
        def iter_content(self,size): yield payload
    def get(url,**kwargs):
        calls.append((url,kwargs))
        return Response()
    import requests
    monkeypatch.setattr(requests,'get',get)
    first,url=locate_source('municipios_anos_iniciais',tmp_path)
    assert first.read_bytes()==payload
    second,_=locate_source('municipios_anos_iniciais',tmp_path)
    assert first==second and len(calls)==1
    assert calls[0][1].get('verify',True) is True
    assert not list(tmp_path.glob('*.part'))
