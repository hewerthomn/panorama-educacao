"""Parser dos cabeçalhos reais, sem usar metas nem recalcular o Ideb."""
import csv
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import re
import unicodedata
from zipfile import ZipFile
from openpyxl import load_workbook
from ..catalog import UFS
from .catalog import STAGES


def normalized(value):
    return ' '.join(unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower().split())


def code(value, length):
    """Excel armazena os códigos como números; identidade sempre textual."""
    try:
        number = Decimal(str(value).strip())
        if not number.is_finite() or number != number.to_integral_value() or number < 0:
            raise ValueError()
        result = str(int(number)).zfill(length)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError(f'Código inválido: {value!r}') from None
    if len(result) != length:
        raise ValueError(f'Código com tamanho inválido: {value!r}')
    return result


def numeric(value, *, ideb=False):
    if value is None or str(value).strip() == '':
        return None, 'missing'
    token = str(value).strip()
    if token == '-' or re.fullmatch(r'(?:ND)?\*+|ND|N/D|N\.D\.', token, re.IGNORECASE):
        return None, 'not_published'
    if normalized(token) in ('na', 'n/a', 'nao se aplica'):
        return None, 'not_applicable'
    try:
        number = Decimal(token.replace(',', '.'))
    except InvalidOperation:
        raise ValueError(f'Valor numérico desconhecido: {value!r}') from None
    if not number.is_finite() or (ideb and not 0 <= number <= 10):
        raise ValueError(f'Valor fora da escala: {value!r}')
    return float(number), 'available'


def workbooks(path):
    path = Path(path)
    if path.suffix.lower() == '.zip':
        with ZipFile(path) as archive:
            # O ZIP oficial repete a mesma tabela em ODS e XLSX. Ler uma vez.
            members = [n for n in archive.namelist() if n.lower().endswith('.xlsx')]
            if not members:
                members = [n for n in archive.namelist() if n.lower().endswith(('.xls', '.csv'))]
            for member in members:
                yield member, archive.read(member)
    else:
        yield path.name, path.read_bytes()


def sheets(name, data):
    if name.lower().endswith('.csv'):
        text = data.decode('utf-8-sig')
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=';,\t')
        yield 'CSV', csv.reader(StringIO(text), dialect)
    elif name.lower().endswith('.xls'):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=data)
        for sheet in workbook.sheets():
            yield sheet.name, (sheet.row_values(i) for i in range(sheet.nrows))
    else:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
        try:
            for sheet in workbook:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            workbook.close()


def stage_for(title):
    title = normalized(title)
    if 'iniciais' in title or '(ai)' in title: return 'anos_iniciais'
    if 'finais' in title or '(af)' in title: return 'anos_finais'
    if 'medio' in title or '(em)' in title: return 'ensino_medio'
    raise ValueError(f'Etapa não identificada: {title}')


def layout(rows, level):
    """A linha técnica é detectada, e os cabeçalhos visuais identificam P por edição."""
    preamble = []
    iterator = iter(rows)
    for row in iterator:
        preamble.append(list(row))
        if any(re.fullmatch(r'VL_OBSERVADO_\d{4}', str(v or '').strip()) for v in row):
            technical = {str(v).strip(): i for i, v in enumerate(row) if v is not None}
            observed = {int(k[-4:]): i for k, i in technical.items() if re.fullmatch(r'VL_OBSERVADO_\d{4}', k)}
            visual = next((r for r in preamble if any(str(v).startswith('Taxa de Aprovação - ') for v in r)), None)
            p_columns = {}
            if visual:
                for i, label in enumerate(visual):
                    if 'Taxa de Aprovação' not in str(label): continue
                    match = re.search(r'(20\d{2})', str(label))
                    end = next((j for j in range(i+1, len(visual)) if visual[j] is not None), len(visual))
                    candidates = {j for r in preamble[:-1] for j in range(i, end)
                                  if j < len(r) and normalized(r[j]) == 'indicador de rendimento (p)'}
                    if match and len(candidates) == 1:
                        p_columns[int(match[1])] = candidates.pop()
            # Algumas UFs não repetem os nomes de identidade na linha técnica.
            if level == 'state':
                if not visual or normalized(visual[0]) != 'regiao/ unidade da federacao' or normalized(visual[1]) != 'rede':
                    raise ValueError('Cabeçalho de UF desconhecido')
                technical.update({'UF_NAME': 0, 'REDE': 1})
            required = ['REDE'] + ([] if level == 'state' else ['SG_UF', 'CO_MUNICIPIO', 'NO_MUNICIPIO'])
            if level == 'school': required += ['ID_ESCOLA', 'NO_ESCOLA']
            if any(k not in technical for k in required):
                raise ValueError(f'Identidade incompleta: {required}')
            return technical, observed, p_columns, iterator, len(preamble)
        if len(preamble) >= 40: break
    raise ValueError('Cabeçalho VL_OBSERVADO_<ano> não encontrado')


def parse(path, *, level, stage=None, years=(2021, 2023, 2025), ufs=None, source_id='source'):
    results, mappings, seen = [], [], {}
    names = {normalized(name): uf for uf, (_, name) in UFS.items()}
    # Abreviações efetivamente publicadas nas três abas estaduais.
    names.update({normalized(k): v for k, v in {'R. G. do Norte': 'RN', 'R. G. do Sul': 'RS', 'M. G. do Sul': 'MS'}.items()})
    for member, data in workbooks(path):
        for title, rows in sheets(member, data):
            actual_stage = stage or stage_for(title)
            if actual_stage not in STAGES: raise ValueError(f'Etapa inválida: {actual_stage}')
            columns, observed, p_columns, body, header_row = layout(rows, level)
            selected = sorted(set(years) & observed.keys())
            if set(years) - observed.keys():
                raise ValueError(f'{member}/{title}: edições ausentes: {sorted(set(years)-observed.keys())}')
            mappings.append({'member': member, 'sheet': title, 'stage': actual_stage, 'header_row': header_row,
                             'columns': {key: value+1 for key, value in columns.items()
                                         if key in ('SG_UF','CO_MUNICIPIO','NO_MUNICIPIO','ID_ESCOLA','NO_ESCOLA','UF_NAME','REDE')
                                         or any(key == f'{prefix}_{y}' for y in selected for prefix in
                                                ('VL_OBSERVADO','VL_NOTA_MEDIA','VL_NOTA_MATEMATICA','VL_NOTA_PORTUGUES'))},
                             'yield_columns': {str(y): p_columns.get(y, columns.get(f'VL_INDICADOR_REND_{y}', -1))+1 for y in selected}})
            for offset, row in enumerate(body, header_row+1):
                def get(key):
                    i = columns.get(key)
                    return row[i] if i is not None and i < len(row) else None
                uf = names.get(normalized(get('UF_NAME'))) if level == 'state' else str(get('SG_UF') or '').strip().upper()
                # Excluir regiões e notas de rodapé; não são observações de UF.
                if uf not in UFS or (ufs and uf not in ufs): continue
                municipality = code(get('CO_MUNICIPIO'), 7) if level != 'state' else None
                school = code(get('ID_ESCOLA'), 8) if level == 'school' else None
                raw_network = str(get('REDE') or '').strip()
                network = re.sub(r'\s*(?:\(\d+\))+\s*$', '', raw_network).strip()
                if not network: raise ValueError(f'Rede ausente em {member}:{offset}')
                for year in selected:
                    value = row[observed[year]]
                    score, status = numeric(value, ideb=True)
                    p = p_columns.get(year, columns.get(f'VL_INDICADOR_REND_{year}'))
                    components = {'yield': numeric(row[p])[0] if p is not None else None}
                    for key, prefix in [('performance','VL_NOTA_MEDIA'),('math','VL_NOTA_MATEMATICA'),('portuguese','VL_NOTA_PORTUGUES')]:
                        components[key] = numeric(get(f'{prefix}_{year}'))[0]
                    record = {'year': year, 'uf': uf, 'uf_code': str(UFS[uf][0]), 'level': level,
                              'municipality_id': municipality, 'municipality_name': get('NO_MUNICIPIO'),
                              'school_id': school, 'name': get('NO_ESCOLA') if school else get('NO_MUNICIPIO') if municipality else UFS[uf][1],
                              'stage': actual_stage, 'network': network, 'network_label': raw_network, 'ideb': score, 'status': status,
                              'raw_value': str(value) if status != 'available' and value is not None else None,
                              'components': components, 'source_id': source_id}
                    key = (year, uf, level, municipality, school, actual_stage, network)
                    if key in seen:
                        if seen[key] != record: raise ValueError(f'Resultados duplicados conflitantes: {key}')
                        continue
                    seen[key] = record
                    results.append(record)
    if not results: raise ValueError(f'Nenhuma observação encontrada em {path}')
    return results, mappings
